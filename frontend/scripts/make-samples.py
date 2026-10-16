"""Write the BOM test files for Step 1's Import BOM tab: public/samples/heatgun-bom.{csv,xlsx,pdf}.

    python scripts/make-samples.py            # needs openpyxl; the PDF needs Chrome or Edge

Same rows as backend/testdata/heatgun_bom.csv (CBP ruling N328954's heat gun),
laid out like a supplier BOM: line, part number and quantity columns (the
importer ignores them), a document header, and the importer's real columns.
The company and part numbers are made up; the file says so.
"""
import csv
import html
import os
import shutil
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "..", "public", "samples")
CAPS_TH = "motor_fan_assembly; pcba; heating_element; housing_moulding; switch_cord_minor_parts; final_assembly"
CAPS_CN = "motor_fan_assembly; pcba; heating_element; housing_moulding; switch_cord_minor_parts"

HEAD = ["Line", "Part no.", "Type", "Name", "Qty", "Country", "Plant", "Group", "Plant capabilities"]
ROWS = [
    ("component", "HG-1001", "main housing", 1, "Thailand", "Rayong plant", "housing_moulding", CAPS_TH),
    ("component", "HG-1002", "front cover", 1, "Thailand", "Rayong plant", "housing_moulding", ""),
    ("component", "HG-2001", "heating element", 1, "Thailand", "Rayong plant", "heating_element", ""),
    ("component", "HG-3001", "electric motor", 1, "China", "China suppliers", "motor_fan_assembly", CAPS_CN),
    ("component", "HG-3002", "fan", 1, "China", "China suppliers", "motor_fan_assembly", ""),
    ("component", "HG-3003", "motor support", 1, "China", "China suppliers", "motor_fan_assembly", ""),
    ("component", "HG-4001", "printed circuit board assembly (PCBA)", 1, "China", "China suppliers", "pcba", ""),
    ("component", "HG-5001", "switch", 1, "China", "China suppliers", "switch_cord_minor_parts", ""),
    ("component", "HG-5002", "power cord", 1, "China", "China suppliers", "switch_cord_minor_parts", ""),
    ("component", "HG-5003", "wiring", 4, "China", "China suppliers", "switch_cord_minor_parts", ""),
    ("component", "HG-5004", "knob", 1, "China", "China suppliers", "switch_cord_minor_parts", ""),
    ("component", "HG-5005", "back cover", 1, "China", "China suppliers", "switch_cord_minor_parts", ""),
    ("component", "HG-5006", "support ring, heat insulator and fasteners", 1, "China", "China suppliers", "switch_cord_minor_parts", ""),
    ("operation", "OP-10", "press fit the fan to the motor spindle and attach the fan cover", "", "Thailand", "Rayong plant", "final_assembly", ""),
    ("operation", "OP-20", "crimp plugs to the heating element wires, add heat shrink tubing, solder wires to the switch and plug the heating element into the switch and PCBA", "", "Thailand", "Rayong plant", "final_assembly", ""),
    ("operation", "OP-30", "fasten the PCBA to the main housing, assemble the temperature knob, connect the PCBA to the motor, add the back cover and motor support", "", "Thailand", "Rayong plant", "final_assembly", ""),
    ("operation", "OP-40", "secure the power cord, add the switch slider, close the housing with screws, test, label and pack", "", "Thailand", "Rayong plant", "final_assembly", ""),
]
TABLE = [[i + 1, r[1], r[0], r[2], r[3], r[4], r[5], r[6], r[7]] for i, r in enumerate(ROWS)]
TITLE = "Bill of materials · 12.5-amp heat gun, two fan speeds, 1500 W"
NOTE = "Example test file for Origin Proof. Rows follow the facts of public CBP ruling N328954; company and part numbers are made up."


def write_csv():
    with open(os.path.join(OUT, "heatgun-bom.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(HEAD)
        w.writerows(TABLE)


def write_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "BOM"
    ws.append(HEAD)  # the importer reads the first sheet's first row as the header
    for row in TABLE:
        ws.append(row)
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="E2EAFA")
    for cell in ws[1]:
        cell.font, cell.fill = bold, fill
    widths = [6, 10, 11, 48, 5, 10, 16, 24, 60]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    for row in ws.iter_rows(min_row=2):
        row[3].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    about = wb.create_sheet("About")
    about.append([TITLE])
    about.append([NOTE])
    wb.save(os.path.join(OUT, "heatgun-bom.xlsx"))


def browser():
    for p in [
        r"C:\Program Files\Google\Chrome\Application\chrome.exe",
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
        "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
    ]:
        if os.path.exists(p):
            return p
    return shutil.which("google-chrome") or shutil.which("chromium") or sys.exit("Chrome or Edge is needed for the PDF")


def write_pdf():
    cells = "".join(
        "<tr>" + "".join(f"<td>{html.escape(str(v))}</td>" for v in row) + "</tr>" for row in TABLE
    )
    doc = f"""<!doctype html><meta charset="utf-8"><style>
@page {{ size: A4 landscape; margin: 14mm; }}
body {{ font: 9pt Helvetica, Arial, sans-serif; color: #13213a; }}
h1 {{ font-size: 14pt; margin: 0 0 2mm; }} p {{ margin: 0 0 1mm; color: #47566f; }}
.meta {{ display: flex; gap: 10mm; margin: 3mm 0 5mm; }} .meta b {{ display: block; color: #13213a; }}
table {{ border-collapse: collapse; width: 100%; }}
th, td {{ border-bottom: 0.4pt solid #c4d0e2; padding: 1.6mm 2mm; text-align: left; vertical-align: top; }}
th {{ background: #e2eafa; font-weight: bold; }}
td:nth-child(4) {{ width: 70mm; }} td:nth-child(9) {{ width: 52mm; }}
footer {{ margin-top: 5mm; font-size: 8pt; color: #47566f; }}
</style>
<h1>{html.escape(TITLE)}</h1>
<p>Northbay Tool Works (example) · Rayong plant, Thailand · Rev. C · 16 Oct 2026</p>
<div class="meta"><div><b>Assembled in</b>Thailand</div><div><b>Plants</b>Rayong plant (TH), China suppliers (CN)</div><div><b>Family</b>motorised tools and small appliances</div></div>
<table><thead><tr>{"".join(f"<th>{h}</th>" for h in HEAD)}</tr></thead><tbody>{cells}</tbody></table>
<footer>{html.escape(NOTE)}</footer>"""
    src = os.path.join(OUT, "_bom.html")
    with open(src, "w", encoding="utf-8") as f:
        f.write(doc)
    pdf = os.path.join(OUT, "heatgun-bom.pdf")
    subprocess.run([browser(), "--headless=new", "--disable-gpu", "--no-pdf-header-footer", f"--print-to-pdf={pdf}", "file:///" + src.replace(os.sep, "/")], check=True, capture_output=True)
    os.remove(src)


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    write_csv()
    write_xlsx()
    write_pdf()
    print("wrote", sorted(os.listdir(OUT)))
